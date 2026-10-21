"""CSV and XLSX exports of the screening lists: method record, round trip, rule identity, range rows and the queue."""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile

import openpyxl
import pytest
from app.ingestion.load_screening_inputs import read_persistence_source
from app.screening_export import GATE_FAILURE_DETAILS
from school_scoring.screening import PACKAGE_VERSION, RULES, GateFailure
from school_scoring.screening.rules import TIE_RULE
from screening_api_helpers import load_synthetic, parse_csv_export
from screening_db_factories import AWKWARD_NAMES as AWK
from screening_db_factories import record, write_persistence_files

RULE_IDS = list(RULES)
CSV = "/api/v1/screening/exports/top.csv"
XLSX = "/api/v1/screening/exports/top.xlsx"
INCOMPLETE = "/api/v1/screening/exports/incomplete.csv"

TOP_RECORD_KEYS = [
    "export",
    "rule_id",
    "rule_label",
    "rule_provenance",
    "rule_provenance_note",
    "policy_role",
    "policy_role_note",
    "score_kind",
    "formula",
    "weights",
    "uses_urgency",
    "input_version",
    "input_sha256",
    "anchors_sha256",
    "normalization_anchors",
    "run_id",
    "package_version",
    "tie_rule",
    "sort_columns",
    "sort_keys",
    "point_rank_scope",
    "rank_bounds_scope",
    "cohort_counts",
    "filters",
    "filters_effect",
    "display_name",
    "csv_text_guard",
    "k",
    "k_meaning",
    "include_ranges",
    "range_schools_reaching_k",
    "range_schools_included",
    "range_school_convention",
    "rows_listed",
    "how_to_rebuild",
    "generated_at",
    "caveat",
]


@pytest.fixture
def loaded(api, connection):
    return load_synthetic(connection)


def export(api, path: str, **params) -> bytes:
    response = api.get(path, params=params)
    assert response.status_code == 200, response.text
    return response.content


def top(api, **params):
    return parse_csv_export(export(api, CSV, **params))


def point_rows(rows):
    return [r for r in rows if r["row_kind"] == "point"]


# ---------------------------------------------------------------------------------------------
# Rebuilding the list.
# ---------------------------------------------------------------------------------------------


@pytest.mark.parametrize("rule_id", RULE_IDS)
def test_the_csv_rebuilds_the_list_from_its_declared_keys(api, loaded, rule_id):  # API7, API12
    method, rows = top(api, rule=rule_id, k=50)

    listed = point_rows(rows)
    sort_columns = method["sort_columns"].split(",")
    assert sort_columns == (
        ["sort_key_primary", "score_point"] if rule_id.startswith("need_first") else ["score_point"]
    )
    # Sort the declared keys descending, then the name ascending, ignoring the order of the file.
    rebuilt = sorted(listed, key=lambda r: (*[-float(r[c]) for c in sort_columns], r["source_school_name"]))
    assert [int(r["rank_point"]) for r in rebuilt] == list(range(1, len(listed) + 1))
    assert rebuilt == listed

    engine = loaded.runs[rule_id]
    assert len(listed) == engine.manifest.point_rank_count
    by_name = {r.school_name: r for r in engine.results}
    assert [(r["source_school_name"], int(r["rank_point"])) for r in listed] == [
        (r.school_name, r.rank_point)
        for r in sorted((r for r in engine.results if r.rank_point), key=lambda r: r.rank_point)
    ]
    for row in listed:
        # Twelve places (R20), enough that rebuilding by these keys gives the engine's order above.
        assert float(row["score_point"]) == round(by_name[row["source_school_name"]].score_point, 12)


def test_k_cuts_the_list_at_a_point_rank_and_a_name_tie_is_broken_by_code_point(api, connection):  # API7
    # Three schools with identical components tie on score and are ordered by name in code point order.
    tied = tuple(record(name) for name in ("beta", "Zulu", "alpha"))
    load_synthetic(connection, tied)

    method, rows = top(api, rule="additive_u", k=2)

    assert [(r["source_school_name"], r["rank_point"]) for r in rows] == [("Zulu", "1"), ("alpha", "2")]
    assert method["k"] == "2" and method["rows_listed"] == "point=2, range=0"


def test_names_are_exported_byte_for_byte(api, loaded):  # API12
    _, rows = top(api, rule="additive_u", k=50)

    assert {r["source_school_name"] for r in rows if r["row_kind"] == "point"} == set(AWK)
    assert (
        len({r["source_school_name"] for r in rows if r["row_kind"] == "point"}) == 9
    )  # both accent forms survive


def test_values_are_the_engine_s_to_twelve_places(api, loaded):  # API12, R20
    for rule_id in ("additive_u", "bottleneck_u", "need_first_no_u"):
        _, rows = top(api, rule=rule_id, k=50)
        engine = {r.school_name: r for r in loaded.runs[rule_id].results}
        for row in point_rows(rows):
            result = engine[row["source_school_name"]]
            for component in "NARQEL":
                assert float(row[component]) == round(getattr(result.components, component), 12)
            for c in result.contributions:
                assert float(row[f"contribution_{c.component.value}_lower"]) == round(c.lower, 12)
                assert float(row[f"contribution_{c.component.value}_upper"]) == round(c.upper, 12)
            if result.maximum:
                assert float(row["maximum_term"]) == round(result.maximum.contribution, 12)
                assert row["maximum_attained_by"] == ";".join(c.value for c in result.maximum.attained_by)
            assert float(row["score_lower"]) == round(result.score_lower, 12) and float(
                row["score_upper"]
            ) == round(result.score_upper, 12)


def test_unknown_values_are_empty_cells_never_zero(api, loaded):  # API12
    _, rows = top(api, rule="additive_u", k=50, include_ranges=True)

    air = next(r for r in rows if r["source_school_name"] == "Interval No Air")
    assert air["Q"] == "" and air["contribution_Q_lower"] != air["contribution_Q_upper"]
    assert air["missing_environmental"] == "Q"
    assert (
        float(air["contribution_Q_lower"]) == 0.0 and float(air["contribution_Q_upper"]) == 0.05
    )  # bounds, not a guess


# ---------------------------------------------------------------------------------------------
# Numbers: up to 12 decimal places, trailing zeros removed, no float noise (R20).
# ---------------------------------------------------------------------------------------------

NUMBER = re.compile(r"-?\d+(\.\d{0,11}[1-9])?")
NUMERIC_COLUMN = re.compile(
    r"[NAURQEL]|sort_key_primary|contribution_\w+|need_part_contribution_\w+|maximum_(component|term)"
    r"|score_(lower|upper|point)|latitude|longitude|rank_(point|best|worst)"
)


def cell(rows, name: str, column: str) -> str:
    return next(r for r in rows if r["source_school_name"] == name)[column]


def numbers(rows) -> list[str]:
    return [text for row in rows for column, text in row.items() if text and NUMERIC_COLUMN.fullmatch(column)]


def xlsx_numbers(content: bytes) -> list[str]:
    """The numbers as the workbook stores them: the text of every value that is not a string."""
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        xml = archive.read("xl/worksheets/sheet2.xml").decode()
    return re.findall(r'<c r="[A-Z]+\d+"(?: s="\d+")? t="n"><v>([^<]*)</v></c>', xml)


def test_a_value_with_float_noise_is_written_clean(api, loaded):  # R20
    _, rows = top(api, rule="additive_u", k=50)

    # The engine's floats are 0.5650000000000001, 0.7150000000000001 and 0.0025000000000000005.
    assert cell(rows, " Leading Space School", "A") == "0.565"
    assert cell(rows, "Trailing Space School ", "N") == "0.715"
    assert cell(rows, "Trailing Space School ", "contribution_L_lower") == "0.0025"


def test_a_value_that_needs_all_twelve_places_keeps_them(api, loaded):  # R20
    _, rows = top(api, rule="additive_u", k=50)

    # 0.08279104477611941 and 0.0029411764705882344, rounded at the twelfth place.
    assert cell(rows, "Trailing Space School ", "contribution_A_lower") == "0.082791044776"
    assert cell(rows, "Trailing Space School ", "contribution_Q_lower") == "0.002941176471"


def test_every_number_in_every_export_has_no_trailing_zero_and_at_most_twelve_places(api, loaded):  # R20
    _, rows = top(api, rule="bottleneck_u", k=50, include_ranges=True)
    assert cell(rows, " Leading Space School", "contribution_U_lower") == "0"  # the engine's 0.0
    _, queue = parse_csv_export(export(api, INCOMPLETE, rule="additive_u"))
    workbook = xlsx_numbers(export(api, XLSX, rule="need_first_u", k=50))

    for name, written in (
        ("top.csv", numbers(rows)),
        ("incomplete.csv", numbers(queue)),
        ("top.xlsx", workbook),
    ):
        assert written, name
        assert [t for t in written if not NUMBER.fullmatch(t)] == [], name


def test_screening_scores_stay_within_0_and_1_after_rounding(api, loaded, connection):  # R20
    # Shares that add to one can add to 1.0000000000000002 in floating point, and a stored score keeps what it was
    # given. These two stand for a score a bit past each end.
    stored = {1: 1.0000000000000002, 2: -1e-12}
    for rank, score in stored.items():
        connection.execute(
            "update screening_results r set score_point = %(s)s, score_lower = %(s)s, score_upper = %(s)s "
            "from screening_runs u where u.run_id = r.run_id and u.rule_id = 'additive_u' and r.rank_point = %(rank)s",
            {"s": score, "rank": rank},
        )

    _, rows = top(api, rule="additive_u", k=2)
    assert [(r["score_point"], r["score_lower"], r["score_upper"]) for r in rows] == [("1",) * 3, ("0",) * 3]
    sheet = openpyxl.load_workbook(io.BytesIO(export(api, XLSX, rule="additive_u", k=2)))["Top schools"]
    header = [c.value for c in sheet[1]]
    scores = [
        [row[header.index(c)] for c in ("score_point", "score_lower", "score_upper")]
        for row in sheet.iter_rows(min_row=2, values_only=True)
    ]
    assert scores == [[1, 1, 1], [0, 0, 0]]


def test_the_method_record_writes_its_shares_and_anchors_the_same_way(api, loaded):  # R20
    method, _ = top(api, rule="additive_u", k=50)
    # The engine's anchors include 135.54999999999998, 0.030000000000000006 and 0.36999999999999994.
    assert "p95=135.55" in method["normalization_anchors"]
    assert "p05=0.03, p95=0.37" in method["normalization_anchors"]
    without, _ = top(api, rule="additive_no_u", k=50)
    assert "N 12/19 (0.631578947368)" in without["weights"]  # 12/19 is 0.631578947368421

    queue_method, _ = parse_csv_export(export(api, INCOMPLETE, rule="additive_no_u"))
    for written in (method, without, queue_method):
        figures = re.findall(r"\d+\.\d+", written["weights"] + " " + written["normalization_anchors"])
        assert figures and [f for f in figures if not NUMBER.fullmatch(f)] == []


# ---------------------------------------------------------------------------------------------
# The method record.
# ---------------------------------------------------------------------------------------------


def test_the_method_record_names_everything_needed_to_read_the_file(api, connection, loaded):  # API17
    before = dt.datetime.now(dt.timezone.utc).replace(microsecond=0)
    method, _ = top(api, rule="additive_no_u")
    after = dt.datetime.now(dt.timezone.utc)

    assert list(method) == TOP_RECORD_KEYS
    rule = RULES["additive_no_u"]
    manifest = loaded.runs["additive_no_u"].manifest
    stored = connection.execute(
        "select run_id, input_sha256, anchors_sha256, package_version from screening_runs where rule_id = 'additive_no_u'"
    ).fetchone()
    assert method["export"] == "top schools"
    assert (method["rule_id"], method["rule_label"]) == ("additive_no_u", "Additive, urgency omitted")
    assert method["rule_provenance"] == "revised_paper"
    assert (
        "revised paper" in method["rule_provenance_note"]
        and "not a result of the circulated paper" in method["rule_provenance_note"]
        and "pending" not in method["rule_provenance_note"]
    )
    assert (method["policy_role"], method["score_kind"]) == ("primary", "additive score")
    assert method["policy_role_note"] == "The revised paper's primary list."
    assert (
        method["formula"]
        == rule.formula
        == "S0 = (0.60N + 0.15A + 0.05(R + Q + E + L)) / 0.95 = (12N + 3A + R + Q + E + L) / 19"
    )
    assert method["weights"].startswith("N 12/19 (0.631578947368)") and "U " not in method["weights"]
    assert method["uses_urgency"] == "no"
    assert method["input_version"] == loaded.input_version
    assert (
        method["run_id"],
        method["input_sha256"],
        method["anchors_sha256"],
        method["package_version"],
    ) == (
        str(stored[0]),
        stored[1],
        stored[2],
        stored[3],
    )
    assert (
        method["anchors_sha256"] == manifest.anchors_sha256 and method["package_version"] == PACKAGE_VERSION
    )
    assert re.fullmatch(r"[0-9a-f]{64}", method["input_sha256"])
    assert (
        "aqi (Q)" in method["normalization_anchors"]
        and "land (L): natural 0..1 fraction" in method["normalization_anchors"]
    )
    assert "urgency" not in method["normalization_anchors"]
    assert method["tie_rule"] == TIE_RULE
    assert method["point_rank_scope"] == "complete under this rule (10)"
    assert method["rank_bounds_scope"] == "eligible (12)"
    assert method["cohort_counts"] == "records=15, complete=10, interval=2, verification_queue=3"
    assert (method["filters"], method["k"], method["include_ranges"]) == ("none", "20", "yes")
    assert (method["range_schools_reaching_k"], method["range_schools_included"]) == ("2", "yes")
    assert "Range schools are interval schools" in method["range_school_convention"]
    assert "lower-score position by default" in method["range_school_convention"]
    generated = dt.datetime.strptime(method["generated_at"], "%Y-%m-%dT%H:%M:%SZ").replace(
        tzinfo=dt.timezone.utc
    )
    assert before <= generated <= after
    assert "does not establish condition accuracy, site safety, intervention benefit" in method["caveat"]
    assert "A zero weight does not mean that no hazard was observed." in method["caveat"]


def test_the_caveat_for_an_urgency_rule_says_urgency_is_unresolved_not_omitted(api, loaded):  # API17
    method, _ = top(api, rule="additive_u")

    assert (
        method["rule_provenance"] == "circulated_paper"
        and method["rule_provenance_note"] == "A rule of the circulated paper."
    )
    assert method["policy_role"] == "urgency_sensitivity"
    assert "restores the supplied urgency scalar" in method["policy_role_note"]
    assert method["uses_urgency"] == "yes"
    assert "Urgency enters this rule" in method["caveat"] and "inadmissible" not in method["caveat"]
    assert (
        method["point_rank_scope"] == "complete under this rule (9)"
        and method["rank_bounds_scope"] == "eligible (11)"
    )


def test_the_record_names_the_experiment_behind_a_rules_persistence_and_its_own_limit(
    api, connection, tmp_path
):  # API17, API35
    reference = tmp_path / "reference"
    sets = [
        read_persistence_source(
            reference, write_persistence_files(reference, {AWK[7]: (1.0, 1.0, 1.0)}), "additive_u"
        ),
        read_persistence_source(
            reference,
            write_persistence_files(reference, {AWK[7]: (0.9, 0.9, 0.9)}, urgency_omitted=True),
            "additive_no_u",
        ),
    ]
    load_synthetic(connection, persistence=sets)

    circulated, _ = top(api, rule="additive_u")
    revised, _ = top(api, rule="additive_no_u")

    for method, source in ((circulated, sets[0]), (revised, sets[1])):
        assert method["persistence_experiment"] == source.experiment_id
        assert method["persistence_source_hash"] == source.source_sha256
        assert "at least 0.8 in each of 3 settings, 5000 draws each" in method["persistence_test"]
    assert circulated["persistence_experiment"] != revised["persistence_experiment"]
    assert "urgency-included experiment" in circulated["persistence_limitation"]
    assert "not an isolated estimate of urgency's effect" in revised["persistence_limitation"]
    keys = list(revised)
    assert keys[keys.index("persistence_experiment") : keys.index("generated_at")] == [
        "persistence_experiment",
        "persistence_source_hash",
        "persistence_test",
        "persistence_limitation",
    ]


def test_a_rule_without_persistence_numbers_has_no_persistence_lines_in_its_record(
    api, connection, tmp_path
):  # API17, API34
    reference = tmp_path / "reference"
    sets = [
        read_persistence_source(
            reference, write_persistence_files(reference, {AWK[7]: (1.0, 1.0, 1.0)}), "additive_u"
        )
    ]
    load_synthetic(connection, persistence=sets)

    for rule_id in RULE_IDS:
        method, _ = top(api, rule=rule_id)
        has = [key for key in method if key.startswith("persistence_")]
        assert bool(has) == (rule_id == "additive_u"), rule_id


@pytest.mark.parametrize("rule_id", RULE_IDS)
def test_every_export_names_the_rule_it_was_asked_for(api, loaded, rule_id):  # API8, API3
    rule = RULES[rule_id]
    response = api.get(CSV, params={"rule": rule_id, "k": 5})
    method, rows = parse_csv_export(response.content)

    assert method["rule_id"] == rule_id and method["rule_label"] == rule.label
    assert method["formula"] == rule.formula and method["policy_role"] == rule.policy_role.value
    assert (
        response.headers["content-disposition"]
        == f'attachment; filename="rise_png_screening_top5_{rule_id}.csv"'
    )
    assert response.headers["x-rise-png-decision-use"] == "research-prototype-only"
    assert {
        (r["rule_id"], r["rule_label"], r["provenance"], r["policy_role"], r["score_kind"]) for r in rows
    } == {(rule_id, rule.label, rule.provenance.value, rule.policy_role.value, rule.score_kind.value)}
    columns = rows[0].keys()
    assert ("contribution_U_lower" in columns) == rule.uses_urgency
    if not rule.uses_urgency:
        assert "U" not in columns  # no column at all, not an empty one
        assert "urgency" not in method["weights"].lower() and "U " not in method["weights"]

    workbook = openpyxl.load_workbook(io.BytesIO(api.get(XLSX, params={"rule": rule_id, "k": 5}).content))
    assert dict(workbook["Method record"].iter_rows(min_row=2, values_only=True))["rule_id"] == rule_id


# ---------------------------------------------------------------------------------------------
# Range schools.
# ---------------------------------------------------------------------------------------------


def test_range_schools_have_their_own_flagged_rows_with_bounds_and_no_point_rank(api, loaded):  # API2, API15
    method, rows = top(api, rule="additive_u", k=20, include_ranges=True)

    assert [r["row_kind"] for r in rows] == ["point"] * 9 + ["range"] * 2
    air, land = rows[9], rows[10]
    assert [r["source_school_name"] for r in (air, land)] == ["Interval No Air", "Interval No Land"]
    for r in (air, land):
        assert r["rank_point"] == "" and r["score_point"] == ""
        assert (r["rank_best"], r["rank_worst"]) == ("5", "11")
        assert (
            r["rank_bounds_scope"] == "eligible (11)"
            and r["rank_point_scope"] == "complete under this rule (9)"
        )
        assert float(r["score_lower"]) < float(r["score_upper"])
    assert (air["missing_environmental"], land["missing_environmental"]) == ("Q", "L")
    assert method["include_ranges"] == "yes" and method["rows_listed"] == "point=9, range=2"
    assert all(r["rank_point"] != "" for r in point_rows(rows))


def test_range_rows_are_left_out_unless_asked_for_and_unless_they_could_reach_the_cut(api, loaded):  # API15
    method, without = top(api, rule="additive_u", k=20, include_ranges=False)
    assert [r["row_kind"] for r in without] == ["point"] * 9
    # The record still says that two range schools reach the cut and that they are left out.
    assert (method["range_schools_reaching_k"], method["range_schools_included"]) == ("2", "no")
    assert method["rows_listed"] == "point=9, range=0"

    # Both interval schools have 5 as their best possible rank.
    _, below_cut = top(api, rule="additive_u", k=4, include_ranges=True)
    assert [r["row_kind"] for r in below_cut] == ["point"] * 4
    method, at_cut = top(api, rule="additive_u", k=5, include_ranges=True)
    assert [r["row_kind"] for r in at_cut] == ["point"] * 5 + ["range"] * 2
    assert method["rows_listed"] == "point=5, range=2"


def test_range_schools_are_included_by_default_and_the_record_counts_those_reaching_k_after_filters(
    api, loaded
):  # API20
    method, rows = top(api, rule="additive_u", k=20)
    assert [r["row_kind"] for r in rows].count("range") == 2
    assert (
        method["include_ranges"],
        method["range_schools_reaching_k"],
        method["range_schools_included"],
    ) == (
        "yes",
        "2",
        "yes",
    )

    # A filter shrinks the count of range schools that reach k, whether or not they are listed.
    for flag, listed in ((True, 1), (False, 0)):
        method, rows = top(api, rule="additive_u", k=20, province="North", include_ranges=flag)
        assert method["range_schools_reaching_k"] == "1"
        assert method["rows_listed"] == f"point=5, range={listed}"
        assert method["range_schools_included"] == ("yes" if flag else "no")

    method, rows = top(api, rule="additive_u", k=4)
    assert (method["range_schools_reaching_k"], method["rows_listed"]) == ("0", "point=4, range=0")


def test_a_complete_school_is_never_listed_as_a_range_row_and_a_rule_without_environment_has_none(
    api, loaded
):  # API2
    _, rows = top(api, rule="benchmark", k=50, include_ranges=True)

    assert {r["row_kind"] for r in rows} == {"point"}
    assert (
        len(rows) == 11
    )  # the benchmark reads no environmental component: the interval schools are complete


# ---------------------------------------------------------------------------------------------
# Filters.
# ---------------------------------------------------------------------------------------------


def test_filters_in_an_export_hide_rows_without_renumbering_ranks(api, loaded):  # API4
    method, rows = top(api, rule="additive_u", k=50, province="North", include_ranges=True)

    assert [(r["source_school_name"], r["rank_point"]) for r in point_rows(rows)] == [
        (AWK[2], "5"),
        (AWK[3], "6"),
        (AWK[1], "7"),
        (AWK[4], "8"),
        (AWK[0], "9"),
    ]
    assert {r["rank_point_scope"] for r in rows} == {
        "complete under this rule (9)"
    }  # the scope is the full run's
    assert [r["source_school_name"] for r in rows if r["row_kind"] == "range"] == ["Interval No Air"]
    assert method["filters"] == "province=North"
    assert "Filters only hide rows" in method["filters_effect"]

    only_district = top(api, rule="additive_u", k=50, province="North", district="North Two")[1]
    assert [(r["source_school_name"], r["rank_point"]) for r in only_district] == [
        (AWK[3], "6"),
        (AWK[4], "8"),
    ]


# ---------------------------------------------------------------------------------------------
# The spreadsheet.
# ---------------------------------------------------------------------------------------------


def test_the_xlsx_holds_a_method_record_sheet_and_the_same_table_as_the_csv(
    api, loaded
):  # API7, API12, API17
    params = {"rule": "bottleneck_no_u", "k": 50, "include_ranges": True}
    csv_method, csv_rows = top(api, **params)
    workbook = openpyxl.load_workbook(io.BytesIO(export(api, XLSX, **params)))

    assert workbook.sheetnames == ["Method record", "Top schools"]
    sheet_method = dict(list(workbook["Method record"].iter_rows(min_row=2, values_only=True)))
    assert next(workbook["Method record"].iter_rows(max_row=1, values_only=True)) == ("key", "value")
    assert list(sheet_method) == list(csv_method)
    assert {k: v for k, v in sheet_method.items() if k != "generated_at"} == {
        k: v for k, v in csv_method.items() if k != "generated_at"
    }

    table = list(workbook["Top schools"].iter_rows(values_only=True))
    header, body = table[0], table[1:]
    assert list(header) == list(csv_rows[0])
    assert len(body) == len(csv_rows) == 12  # 10 complete under the rule and 2 interval schools
    for csv_row, cells in zip(csv_rows, body, strict=True):
        for column, cell in zip(header, cells, strict=True):
            text = csv_row[column]
            if text == "":
                assert cell is None, column
            elif isinstance(cell, (int, float)):
                assert float(text) == cell, column  # the same numbers, to the same twelve places
            else:
                assert cell == text, column


def test_a_school_name_that_looks_like_a_formula_stays_text_in_the_xlsx(api, connection):  # API19
    risky = ("=1+1 School", "+Plus School", "@At School", "-Minus School")
    load_synthetic(connection, tuple(record(name) for name in risky))

    content = export(api, XLSX, rule="additive_u", k=10)

    sheet = openpyxl.load_workbook(io.BytesIO(content))["Top schools"]
    header = [c.value for c in sheet[1]]
    column = header.index("source_school_name")
    cells = [row[column] for row in sheet.iter_rows(min_row=2)]
    assert sorted(c.value for c in cells) == sorted(risky)
    assert {c.data_type for c in cells} == {"s"}
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        data_sheets = [archive.read(n).decode() for n in archive.namelist() if n.startswith("xl/worksheets/")]
    assert all("<f>" not in xml for xml in data_sheets)


# ---------------------------------------------------------------------------------------------
# The queue.
# ---------------------------------------------------------------------------------------------


def test_the_incomplete_export_lists_exactly_the_queue_with_missing_details(api, loaded):  # API16
    method, rows = parse_csv_export(export(api, INCOMPLETE, rule="additive_u"))

    assert [r["source_school_name"] for r in rows] == [
        "Éclair queue",
        "Zed queue",
        "alpha queue",
        "Beta queue",
    ]
    assert {r["row_kind"] for r in rows} == {"queue"} and {r["cohort"] for r in rows} == {
        "verification_queue"
    }
    by_name = {r["source_school_name"]: r for r in rows}
    assert by_name["Zed queue"]["gate_failures"] == "missing_stock_field"
    assert by_name["Zed queue"]["missing_input_fields"] == "teachers"
    assert by_name["Zed queue"]["gate_failure_details"] == (
        "A stock count, facility count, power source or water source is missing."
    )
    assert by_name["Éclair queue"]["gate_failures"] == "urgency_unknown"
    assert by_name["Éclair queue"]["missing_input_fields"] == "urgency_source_scalar"
    assert (
        by_name["Éclair queue"]["N"] != "" and by_name["Éclair queue"]["U"] == ""
    )  # what is known stays visible
    # An unrecognized value is a gate failure, but the input is not missing.
    assert by_name["alpha queue"]["gate_failures"] == "unrecognized_locality"
    assert by_name["alpha queue"]["missing_input_fields"] == ""
    assert {"teachers", "urgency_source_scalar", "aqi_p95_selected"} <= set(
        by_name["Beta queue"]["missing_input_fields"].split(";")
    )
    assert method["export"] == "schools held for verification" and method["rows_listed"] == "queue=4"
    assert method["rule_id"] == "additive_u" and method["input_sha256"]
    assert all(r["gate_failure_details"] for r in rows)


def test_the_queue_depends_on_the_rule_and_on_the_filters(api, loaded):  # API16, API4
    _, no_urgency = parse_csv_export(export(api, INCOMPLETE, rule="additive_no_u"))
    assert [r["source_school_name"] for r in no_urgency] == ["Zed queue", "alpha queue", "Beta queue"]

    method, north = parse_csv_export(export(api, INCOMPLETE, rule="additive_u", province="North"))
    assert [r["source_school_name"] for r in north] == ["Éclair queue", "Zed queue", "alpha queue"]
    assert method["filters"] == "province=North" and method["rows_listed"] == "queue=3"


def test_every_gate_failure_has_a_plain_description():  # API16
    assert set(GATE_FAILURE_DETAILS) == set(GateFailure)


def test_control_characters_in_stored_text_do_not_break_the_xlsx(api, connection):  # API23
    names = ("Bell\x07 School", "Esc\x1b School", "Vertical\x0bTab School")
    records = tuple(record(name, province="Prov\x01ince") for name in names)
    load_synthetic(connection, records)

    workbook = export(api, XLSX, rule="additive_u", k=10)
    _, csv_rows = top(api, rule="additive_u", k=10)

    sheet = openpyxl.load_workbook(io.BytesIO(workbook))["Top schools"]
    header = [c.value for c in sheet[1]]
    stored = {row[header.index("source_school_name")] for row in sheet.iter_rows(min_row=2, values_only=True)}
    # XML cannot hold these characters, so the workbook carries Excel's escape for each.
    assert stored == {"Bell_x0007_ School", "Esc_x001B_ School", "Vertical_x000B_Tab School"}
    provinces = {row[header.index("province")] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert provinces == {"Prov_x0001_ince"}
    # The CSV keeps the characters as they are stored.
    assert {r["source_school_name"] for r in csv_rows} == set(names)


FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def test_the_csv_guards_every_text_cell_that_starts_with_a_formula_character(api, connection):  # API24
    risky = ("=1+1 School", "+Plus School", "-Minus School", "@At School", "\tTab School", "\rReturn School")
    safe = ("Plain School", "'Quoted School", "Mid=dle School")
    load_synthetic(connection, tuple(record(name) for name in (*risky, *safe)))

    method, rows = top(api, rule="additive_u", k=50)
    incomplete_method, _ = parse_csv_export(export(api, INCOMPLETE, rule="additive_u"))

    # In the CSV the name is guarded in both columns; every other name is as stored.
    guarded = {r["school_key"]: (r["source_school_name"], r["display_name"]) for r in rows}
    assert sorted(guarded.values()) == sorted(
        [("'" + name, "'" + name) for name in risky] + [(name, name) for name in safe]
    )
    assert list(rows[0]).index("display_name") == list(rows[0]).index("source_school_name") + 1
    assert "apostrophe" in method["display_name"] and "csv_text_guard" in method["display_name"]
    assert (
        "csv_text_guard" in incomplete_method
        and incomplete_method["csv_text_guard"] == method["csv_text_guard"]
    )

    # The XLSX stores text as text, so its source name stays exact and its display name keeps the apostrophe.
    sheet = openpyxl.load_workbook(io.BytesIO(export(api, XLSX, rule="additive_u", k=50)))["Top schools"]
    header = [c.value for c in sheet[1]]
    in_xlsx = {
        row[header.index("source_school_name")].replace("_x000D_", "\r"): row[
            header.index("display_name")
        ].replace("_x000D_", "\r")
        for row in sheet.iter_rows(min_row=2, values_only=True)
    }
    assert in_xlsx == {**{name: "'" + name for name in risky}, **{name: name for name in safe}}


@pytest.mark.parametrize("start", FORMULA_STARTS)
def test_province_district_and_locality_are_guarded_in_the_csv_but_numbers_are_not(
    api, connection, start
):  # API24
    place = f"{start}CMD(1)"
    # A locality the engine does not recognise holds a school back, so that cell is read from the queue export.
    complete = tuple(record(f"School {i}", province=place, district=place) for i in range(3))
    held = tuple(record(f"Held {i}", province=place, district=place, locality=place) for i in range(2))
    load_synthetic(connection, (*complete, *held))

    _, listed = parse_csv_export(export(api, CSV, rule="additive_u", k=10))
    _, queue = parse_csv_export(export(api, INCOMPLETE, rule="additive_u"))
    assert len(listed) == 3 and len(queue) == 2
    for row in (*listed, *queue):
        assert row["province"] == "'" + place and row["district"] == "'" + place
    for row in queue:
        assert row["locality"] == "'" + place

    # A negative number is a number, so it carries no apostrophe, and every text cell that starts with a formula
    # character is guarded: nothing in the file is left for a spreadsheet to read as a formula.
    def is_number(cell: str) -> bool:
        try:
            float(cell)
        except ValueError:
            return False
        return True

    text = export(api, CSV, rule="additive_u", k=10).decode("utf-8-sig")
    parsed = list(csv.reader(io.StringIO(text, newline="")))
    first = next(i for i, row in enumerate(parsed) if row and row[0] == "row_kind")
    for row in parsed[first + 1 :]:
        for cell in row:
            assert not cell.startswith(FORMULA_STARTS) or is_number(cell), cell


def test_a_negative_number_in_the_csv_is_not_guarded(api, connection):  # API24
    load_synthetic(connection, tuple(record(f"School {i}") for i in range(3)))
    # The latitude of a school is negative, and it is written as a number.
    _, rows = top(api, rule="additive_u", k=10)

    latitudes = [r["latitude"] for r in rows if r["latitude"]]
    assert latitudes and all(v.startswith("-") and float(v) < 0 for v in latitudes)
    assert all(not v.startswith("'") for v in latitudes)


def test_the_csv_opens_in_excel_with_accents_and_one_column_per_record_line(api, loaded):  # API29
    raw = export(api, CSV, rule="additive_no_u", k=5)

    assert raw.startswith(b"\xef\xbb\xbf")  # a byte order mark, so that Excel reads UTF-8
    text = raw.decode("utf-8")
    lines = text.lstrip("\ufeff").split("\n")
    first_data = next(i for i, line in enumerate(lines) if line.startswith("row_kind,"))
    record_lines = lines[:first_data]
    # Every record line is one quoted cell, even when its value holds commas, quotes or semicolons.
    assert all(line.startswith('"# ') and line.endswith('"') for line in record_lines)
    assert any("," in line for line in record_lines)  # for example the weights
    parsed = list(csv.reader(io.StringIO(text.lstrip("\ufeff"), newline="")))
    assert [len(row) for row in parsed[:first_data]] == [1] * first_data
    header_width = len(parsed[first_data])
    assert header_width > 20 and {len(row) for row in parsed[first_data:]} == {header_width}
    # A name with an accent survives the round trip.
    _, queue = parse_csv_export(export(api, INCOMPLETE, rule="additive_u"))
    assert "Éclair queue" in {r["source_school_name"] for r in queue}


def test_urgency_has_no_column_and_no_missing_input_under_an_urgency_omitted_rule(api, loaded):  # API27
    _, with_urgency = parse_csv_export(export(api, INCOMPLETE, rule="additive_u"))
    _, without = parse_csv_export(export(api, INCOMPLETE, rule="additive_no_u"))

    assert "U" in with_urgency[0] and "U" not in without[0]
    beta_with = next(r for r in with_urgency if r["source_school_name"] == "Beta queue")
    beta_without = next(r for r in without if r["source_school_name"] == "Beta queue")
    assert "urgency_source_scalar" in beta_with["missing_input_fields"].split(";")
    # Under the omitted rule the urgency source value is neither read nor asked for.
    assert "urgency_source_scalar" not in beta_without["missing_input_fields"].split(";")
    assert "teachers" in beta_without["missing_input_fields"].split(";")
    _, top_rows = top(api, rule="need_first_no_u", k=50)
    assert "U" not in top_rows[0] and not [c for c in top_rows[0] if c.endswith("_U_lower") or "_U_" in c]
    _, top_with = top(api, rule="need_first_u", k=50)
    assert "U" in top_with[0]
