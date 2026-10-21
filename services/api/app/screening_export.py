"""CSV and XLSX exports of screening lists, each with the method record needed to read and rebuild it.

An export is a table plus a method record.
The CSV starts with a byte order mark, then the record as quoted one-cell lines that start with ``# ``, then one header row and the data rows.
Parse it as CSV (decode with ``utf-8-sig``) and skip the leading one-cell rows that start with ``# ``.
The XLSX holds the record on its own first sheet.
A number is written to at most 12 decimal places with trailing zeros removed, so float noise never reaches the
file (``1``, not ``1.0000000000000002``), and a screening score stays within 0 to 1 after rounding.
A missing value is an empty cell, never zero.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import math
import re
import threading
from collections.abc import Sequence
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import _writer as openpyxl_cell_writer
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from school_scoring.screening import (
    INPUT_CONTRACT,
    Component,
    GateFailure,
    PolicyRole,
    Provenance,
    SchoolRecord,
)

from . import screening_store as store
from .models.screening import PERSISTENCE_LIMITATIONS, AnchorOut, ManifestOut, WeightOut
from .screening_lists import Filters, TopSelection
from .screening_queries import SchoolRow

Cell = str | int | float | None

GATE_FAILURE_DETAILS: dict[GateFailure, str] = {
    GateFailure.MISSING_STOCK_FIELD: "A stock count, facility count, power source or water source is missing.",
    GateFailure.NEGATIVE_STOCK_VALUE: "A stock count is negative.",
    GateFailure.NONPOSITIVE_DENOMINATOR: "Total classrooms or available teachers is zero or negative.",
    GateFailure.UNRECOGNIZED_SERVICE_CODING: "The power or water source text matches no recognized category.",
    GateFailure.UNRECOGNIZED_LOCALITY: "The locality category is missing or not recognized.",
    GateFailure.UNRESOLVED_AREA_CONTEXT: "A district-level input of area context is missing or invalid.",
    GateFailure.UNRESOLVED_REACH: "The walking catchment population is missing or negative.",
    GateFailure.URGENCY_UNKNOWN: "This rule reads urgency and the urgency source value is missing.",
}

PROVENANCE_NOTES: dict[Provenance, str] = {
    Provenance.CIRCULATED_PAPER: "A rule of the circulated paper.",
    Provenance.REVISED_PAPER: (
        "A rule of the revised paper, which leaves urgency out of every primary calculation. "
        "It is not a result of the circulated paper."
    ),
}

POLICY_ROLE_NOTES: dict[PolicyRole, str] = {
    PolicyRole.PRIMARY: "The revised paper's primary list.",
    PolicyRole.STRICT_ALTERNATIVE: "A stricter ordering that puts school need before every other input.",
    PolicyRole.POLICY_SENSITIVITY: "A sensitivity that lets the largest of several inputs lift a school.",
    PolicyRole.URGENCY_SENSITIVITY: (
        "A sensitivity that restores the supplied urgency scalar. It is kept for comparison and is not the primary list."
    ),
    PolicyRole.CONVENTIONAL_REFERENCE: "A conventional benchmark kept for comparison.",
}

CAVEAT = (
    "Research prototype for screening. The list orders existing school records for closer inspection. "
    "It does not establish condition accuracy, site safety, intervention benefit or an optimal investment portfolio. "
    "A missing input is unknown, not zero. "
    "Normalized AQI, reported greenhouse-gas emissions and modal land-cover classes are different contextual measures, "
    "and together they are not a validated health or hazard scale. "
    "Ranks have different scopes: read each with its scope label."
)
CAVEAT_URGENCY_USED = (
    " Urgency enters this rule; how its source value was extracted and what it means are unresolved."
)
CAVEAT_URGENCY_OMITTED = (
    " Urgency is omitted because its primary measurement is inadmissible. "
    "A zero weight does not mean that no hazard was observed."
)

PLACEMENT_CONVENTION = (
    "Range schools are interval schools: an environmental input is missing, so they have a score range and a rank range "
    "and no point rank or point score. "
    "They appear in their own rows (row_kind=range), after the point rows, ordered by best possible rank. "
    "They are not placed in the point order. "
    "The JSON list places them at their lower-score position by default, and at their upper-score position "
    "only when placement=upper is asked for as a labelled sensitivity."
)

TOP_KEYS_NOTE = (
    "k cuts the rule's full list: point rows have rank_point <= k, and range rows have a best possible rank "
    "(rank_best) <= k. Filters then hide rows; they never renumber a rank."
)

TIME_FORMAT = "%Y-%m-%dT%H:%M:%SZ"


DECIMALS = 12
SCORE_COLUMNS = frozenset({"score_lower", "score_upper", "score_point"})


def export_number(value: float, *, score: bool = False) -> float:
    """A float as the exports write it: rounded at the twelfth decimal place, a screening score kept within 0 to 1."""
    if not math.isfinite(value):
        return value
    rounded = round(value, DECIMALS)
    if score:
        rounded = min(1.0, max(0.0, rounded))
    return rounded + 0.0  # adding zero turns a negative zero positive


def number_text(value: float) -> str:
    """A float's text in an export: up to 12 decimal places, trailing zeros removed, so ``0.565`` and never
    ``0.5650000000000001``, and ``1`` and never ``1.0`` or ``1.0000000000000002``."""
    if not math.isfinite(value):
        return repr(value)
    text = f"{export_number(value):.{DECIMALS}f}".rstrip("0").rstrip(".")
    return "0" if text == "-0" else text


def _clean_rows(table: Table) -> list[list[Cell]]:
    """The table's rows with every float as the exports write it."""
    scores = [column in SCORE_COLUMNS for column in table.columns]
    return [
        [
            export_number(v, score=score) if isinstance(v, float) else v
            for v, score in zip(row, scores, strict=True)
        ]
        for row in table.rows
    ]


def _format_weight(weight: WeightOut) -> str:
    return f"{weight.component.value} {weight.fraction} ({number_text(weight.value)})"


def _format_anchor(anchor: AnchorOut) -> str:
    bounds = (
        "natural 0..1 fraction"
        if anchor.p05 is None
        else f"p05={number_text(anchor.p05)}, p95={number_text(anchor.p95)}"
    )
    return f"{anchor.name} ({anchor.component.value}): {anchor.transformation}; {bounds}; valid {anchor.n_valid} of {anchor.n_records}"


FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def looks_like_formula(text: str) -> bool:
    """Whether a spreadsheet could read the text as a formula: its first non-blank character is one of ``FORMULA_STARTS``.

    Leading whitespace, including line breaks and the no-break space, does not hide the character from the application.
    """
    return text.startswith(FORMULA_STARTS) or text.lstrip().startswith(FORMULA_STARTS)


DISPLAY_NAME_NOTE = (
    "display_name is source_school_name with an apostrophe in front when the name starts with =, +, -, @, "
    "a tab or a carriage return, ignoring any leading whitespace, so that a spreadsheet application does not read it as a formula. "
    "In the XLSX source_school_name is exact: use it to join and to rebuild the list. "
    "In the CSV it is guarded like every text cell (see csv_text_guard): join on school_key instead."
)

CSV_TEXT_GUARD_NOTE = (
    "In the CSV every text cell that starts with =, +, -, @, a tab or a carriage return, ignoring any leading whitespace, "
    "has an apostrophe in front, so that a spreadsheet application does not read it as a formula. "
    "That covers source_school_name, province, district and locality as well as display_name. "
    "Numbers are not changed, and a negative number stays a number. "
    "To recover a value, remove one leading apostrophe when the text after it, ignoring leading whitespace, starts with one of those; "
    "a value that was stored with an apostrophe there already is ambiguous, so join on school_key. "
    "The XLSX stores every text cell as text and is not changed."
)


def display_name(source_school_name: str) -> str:
    return "'" + source_school_name if looks_like_formula(source_school_name) else source_school_name


def method_record(
    *,
    export: str,
    manifest: ManifestOut,
    run: store.StoredRun,
    version: store.InputVersion,
    filters: Filters,
    generated_at: dt.datetime,
    top: TopSelection | None = None,
    queue_count: int | None = None,
    persistence: store.PersistenceSet | None = None,
) -> list[tuple[str, str]]:
    """The facts a reader needs to interpret an export and rebuild its list, as ordered key/value pairs."""
    sort_keys = ", ".join(manifest.sort_columns)
    record: list[tuple[str, str]] = [
        ("export", export),
        ("rule_id", manifest.rule_id),
        ("rule_label", manifest.rule_label),
        ("rule_provenance", manifest.provenance.value),
        ("rule_provenance_note", PROVENANCE_NOTES[manifest.provenance]),
        ("policy_role", manifest.policy_role.value),
        ("policy_role_note", POLICY_ROLE_NOTES[manifest.policy_role]),
        ("score_kind", manifest.score_kind),
        ("formula", manifest.rule_formula),
        ("weights", "; ".join(_format_weight(w) for w in manifest.weights)),
        ("uses_urgency", "yes" if manifest.uses_urgency else "no"),
        ("input_version", version.input_version),
        ("input_sha256", run.input_sha256),
        ("anchors_sha256", run.anchors_sha256),
        ("normalization_anchors", " | ".join(_format_anchor(a) for a in manifest.anchors)),
        ("run_id", str(run.run_id)),
        ("package_version", run.package_version),
        ("tie_rule", manifest.tie_rule),
        ("sort_columns", ",".join(manifest.sort_columns)),
        ("sort_keys", f"{sort_keys} descending, then source_school_name ascending"),
        ("point_rank_scope", manifest.point_rank_scope),
        ("rank_bounds_scope", manifest.rank_bounds_scope),
        ("cohort_counts", ", ".join(f"{name}={count}" for name, count in manifest.cohort_counts.items())),
        ("filters", filters.describe()),
        ("filters_effect", "Filters only hide rows. Ranks, bounds and anchors are those of the full run."),
        ("display_name", DISPLAY_NAME_NOTE),
        ("csv_text_guard", CSV_TEXT_GUARD_NOTE),
    ]
    if top is not None:
        record += [
            ("k", str(top.k)),
            ("k_meaning", TOP_KEYS_NOTE),
            ("include_ranges", "yes" if top.include_ranges else "no"),
            ("range_schools_reaching_k", str(len(top.ranges_reaching_k))),
            ("range_schools_included", "yes" if top.include_ranges else "no"),
            ("range_school_convention", PLACEMENT_CONVENTION),
            ("rows_listed", f"point={len(top.point_rows)}, range={len(top.range_rows)}"),
            (
                "how_to_rebuild",
                "Keep rows with row_kind=point. Sort by "
                f"{sort_keys} descending, then source_school_name ascending in Unicode code point order. "
                "The result lists the rows in rank_point order.",
            ),
        ]
    if queue_count is not None:
        record += [
            (
                "queue_definition",
                "Schools with a gate failure under this rule, or whose rule reads urgency and urgency is unknown.",
            ),
            ("rows_listed", f"queue={queue_count}"),
        ]
    if persistence is not None:
        record += [
            ("persistence_experiment", persistence.experiment_id),
            # Not named ``..._sha256``: a client counts the record's ``..._sha256`` lines as the source files of the list,
            # and the persistence table is evidence for a flag, not an input of the ranking.
            ("persistence_source_hash", persistence.source_sha256),
            (
                "persistence_test",
                f"Lower top-20 frequency at least {number_text(persistence.threshold)} in each of "
                f"{len(persistence.settings)} settings, {persistence.draws} draws each. "
                "A frequency is a share of draws under declared scenarios, not a probability that a school needs help.",
            ),
            ("persistence_limitation", PERSISTENCE_LIMITATIONS[persistence.rule_id]),
        ]
    record += [
        ("generated_at", generated_at.astimezone(dt.timezone.utc).strftime(TIME_FORMAT)),
        (
            "caveat",
            CAVEAT + (CAVEAT_URGENCY_USED if manifest.uses_urgency else CAVEAT_URGENCY_OMITTED),
        ),
    ]
    return record


# ---------------------------------------------------------------------------------------------
# Tables.
# ---------------------------------------------------------------------------------------------

COMPONENT_COLUMNS = tuple(c.value for c in Component)
NEED_PART_NAMES = (
    "permanent_classroom_deficit",
    "teacher_housing_deficit",
    "facility_presence_deficit",
    "service_deficit",
    "locality_disadvantage",
)


@dataclass(frozen=True)
class Table:
    columns: tuple[str, ...]
    rows: tuple[tuple[Cell, ...], ...]


def _identity_columns() -> tuple[str, ...]:
    return (
        "row_kind", "rule_id", "rule_label", "provenance", "policy_role", "score_kind",
        "input_version", "input_sha256", "anchors_sha256", "package_version",
        "school_key", "school_id", "source_school_name", "display_name", "province", "district", "locality", "latitude", "longitude",
        "cohort", "gate_failures", "missing_environmental",
    )  # fmt: skip


def _identity_values(kind: str, row: SchoolRow, manifest: ManifestOut, input_version: str) -> list[Cell]:
    i, r = row.identity, row.result
    components = r.components
    return [
        kind, manifest.rule_id, manifest.rule_label, manifest.provenance.value, manifest.policy_role.value,
        manifest.score_kind, input_version, manifest.input_sha256, manifest.anchors_sha256, manifest.package_version,
        i.school_key, None if i.school_id is None else str(i.school_id), i.source_school_name, display_name(i.source_school_name), i.province, i.district, i.locality,
        i.latitude, i.longitude, r.cohort.value, ";".join(f.value for f in r.gate_failures),
        ";".join(c for c in ("Q", "E", "L") if components[c] is None),
    ]  # fmt: skip


def component_columns(manifest: ManifestOut) -> tuple[str, ...]:
    """The component columns of a rule. A rule that does not read urgency has no ``U`` column, not an empty one."""
    return tuple(c for c in COMPONENT_COLUMNS if c != "U" or manifest.uses_urgency)


def top_table(selection: TopSelection, manifest: ManifestOut, input_version: str) -> Table:
    """Point rows, then range rows. Range rows carry bounds only: no point score and no point rank."""
    contribution_columns = tuple(
        f"contribution_{w.component.value}_{side}" for w in manifest.weights for side in ("lower", "upper")
    )
    has_need_parts = manifest.primary_sort_key is None
    has_maximum = bool(manifest.maximum_components)
    columns = (
        *_identity_columns(),
        *component_columns(manifest),
        "sort_key_primary",
        *contribution_columns,
        *(f"need_part_contribution_{name}" for name in NEED_PART_NAMES if has_need_parts),
        *(("maximum_component", "maximum_attained_by", "maximum_term") if has_maximum else ()),
        "score_lower", "score_upper", "score_point",
        "rank_point", "rank_point_scope", "rank_best", "rank_worst", "rank_bounds_scope",
    )  # fmt: skip

    def values(kind: str, row: SchoolRow) -> tuple[Cell, ...]:
        r = row.result
        by_component = {c["component"]: c for c in r.contributions}
        cells = [*_identity_values(kind, row, manifest, input_version)]
        cells += [r.components[c] for c in component_columns(manifest)]
        cells.append(r.sort_key_primary)
        for w in manifest.weights:
            contribution = by_component.get(w.component.value)
            cells += [None, None] if contribution is None else [contribution["lower"], contribution["upper"]]
        if has_need_parts:
            cells += [
                None if r.need_part_contributions is None else r.need_part_contributions[name]
                for name in NEED_PART_NAMES
            ]
        if has_maximum:
            m = r.maximum_term
            cells += (
                [None, None, None]
                if m is None
                else [m["value"], ";".join(m["attained_by"]), m["contribution"]]
            )
        cells += [
            r.score_lower, r.score_upper, r.score_point,
            r.rank_point, r.rank_point_scope,
            r.rank_best, r.rank_worst, r.rank_bounds_scope,
        ]  # fmt: skip
        return tuple(cells)

    rows = [values("point", r) for r in selection.point_rows] + [
        values("range", r) for r in selection.range_rows
    ]
    return Table(columns, tuple(rows))


def missing_input_fields(record: SchoolRecord, uses_urgency: bool) -> tuple[str, ...]:
    """Names of the input-contract fields that are unknown for a school.

    A rule that does not read urgency does not ask for its source value.
    """
    return tuple(
        f.name
        for f in INPUT_CONTRACT
        if getattr(record, f.name) is None and (uses_urgency or f.name != "urgency_source_scalar")
    )


def incomplete_table(
    rows: Sequence[SchoolRow], records: dict[str, SchoolRecord], manifest: ManifestOut, input_version: str
) -> Table:
    """The verification queue: why each school was held back and which inputs are unknown."""
    columns = (
        *_identity_columns(),
        "gate_failure_details",
        "missing_input_fields",
        *component_columns(manifest),
        "rank_point_scope",
        "rank_bounds_scope",
    )
    body = []
    for row in rows:
        r = row.result
        body.append(
            (
                *_identity_values("queue", row, manifest, input_version),
                " ".join(GATE_FAILURE_DETAILS[f] for f in r.gate_failures),
                ";".join(missing_input_fields(records[row.name], manifest.uses_urgency)),
                *(r.components[c] for c in component_columns(manifest)),
                r.rank_point_scope,
                r.rank_bounds_scope,
            )
        )
    return Table(columns, tuple(body))


# ---------------------------------------------------------------------------------------------
# Rendering.
# ---------------------------------------------------------------------------------------------


def _csv_cell(value: Cell) -> str | int:
    """A cell as the CSV writes it. A text cell that a spreadsheet could read as a formula gets an apostrophe in front."""
    if value is None:
        return ""
    if isinstance(value, float):
        return number_text(value)
    if isinstance(value, str) and looks_like_formula(value):
        return "'" + value
    return value


def render_csv(record: Sequence[tuple[str, str]], table: Table) -> bytes:
    """A UTF-8 file with a byte order mark, so that Excel reads accents correctly.

    The method record comes first, one quoted cell per line, ``"# key: value"``, so that Excel keeps each line in one column.
    Then comes the header row and the data rows, every float as ``number_text`` writes it and every text cell that
    starts with a formula character guarded by an apostrophe (see ``CSV_TEXT_GUARD_NOTE``).
    """
    buffer = io.StringIO()
    record_writer = csv.writer(buffer, lineterminator="\n", quoting=csv.QUOTE_ALL)
    for key, value in record:
        record_writer.writerow([f"# {key}: {' '.join(str(value).split())}"])
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(table.columns)
    for row in _clean_rows(table):
        writer.writerow([_csv_cell(v) for v in row])
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


_xlsx_lock = threading.Lock()


def _exact_number_text(value: Any) -> str:
    """openpyxl prints floats with 16 significant digits; the workbook stores the same text as the CSV instead."""
    if isinstance(value, float) and math.isfinite(value):
        return number_text(value)
    return _openpyxl_safe_string(value)


_openpyxl_safe_string = openpyxl_cell_writer.safe_string


# Characters XML cannot hold, plus the carriage return, which an XML reader would turn into a line feed.
_XLSX_ESCAPED = re.compile(f"{ILLEGAL_CHARACTERS_RE.pattern}|\r")


def _xlsx_text(value: str) -> str:
    """These characters are written as Excel's own escape, for example ``_x0007_``, which Excel reads back."""
    return _XLSX_ESCAPED.sub(lambda match: f"_x{ord(match.group()):04X}_", value)


def render_xlsx(record: Sequence[tuple[str, str]], data_sheet: str, table: Table) -> bytes:
    """A workbook with the sheet "Method record" first and the data second. Text is always stored as text."""
    workbook = Workbook()
    method = workbook.active
    method.title = "Method record"
    data = workbook.create_sheet(data_sheet)
    for sheet, rows in ((method, [("key", "value"), *record]), (data, [table.columns, *_clean_rows(table)])):
        for r, values in enumerate(rows, start=1):
            for c, value in enumerate(values, start=1):
                cell = sheet.cell(
                    row=r, column=c, value=_xlsx_text(value) if isinstance(value, str) else value
                )
                if isinstance(value, str):
                    cell.data_type = "s"  # a name that starts with "=" must not become a formula
    method.column_dimensions["A"].width = 26
    method.column_dimensions["B"].width = 120
    data.freeze_panes = "A2"
    buffer = io.BytesIO()
    # The writer's number formatter is module state, so swap it only while this workbook is written.
    with _xlsx_lock:
        openpyxl_cell_writer.safe_string = _exact_number_text
        try:
            workbook.save(buffer)
        finally:
            openpyxl_cell_writer.safe_string = _openpyxl_safe_string
    return buffer.getvalue()
